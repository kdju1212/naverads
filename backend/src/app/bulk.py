"""대량 작업 엔진: 미리보기(검증) → 실행(권한 재검증) → 행별 결과 + 감사 로그.
모든 광고 데이터 쓰기는 이 모듈만 사용한다(확장소재 생성·삭제도 여기서 권한 확인·감사 로그).

광고 데이터 사본은 서버에 없다: 화면(브라우저 보관본)이 변경마다 광고주 계정·이전 값·유형을 같이 보낸다.
서버는 그 계정의 수정 권한만 믿을 근거로 삼고, 실제 대상은 그 계정 API(X-Customer)로만 바꾸므로
다른 광고주의 ID를 섞어 보내도 네이버가 거부한다."""
import io
import json
import os
from datetime import date, datetime, timedelta, timezone

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select, update
from sqlalchemy.orm import Session

from .access import accessible_account_ids
from .models import KST, AccountGoal, AdAccount, AuditLog, AutoBidRule, BulkJob, BulkJobRow, Membership, NaverApiCredential, WorkspaceAdAccount, now
from .naver import NaverApiError, ad_row, client_for, extension_row
from .security import decrypt, rotated

ALLOWED_FIELDS = {
    # 캠페인 create: naver_id는 유형(WEB_SITE|SHOPPING), value는 "이름|일예산"
    # delete(캠페인·광고그룹)는 안의 광고그룹·키워드·소재까지 지워지므로 실행 시 OFF인 것만 지운다
    "campaign": {"status", "daily_budget", "create", "delete"},
    # 광고그룹 create: naver_id는 캠페인, value는 "이름|비즈채널ID|입찰가|하루예산"(입찰가·하루예산은 비워도 됨)
    "adgroup": {"status", "bid_amt", "create", "delete"},
    "keyword": {"status", "bid_amt", "create", "delete"},  # create: naver_id는 광고그룹, value는 "키워드" 또는 "키워드|입찰가"
    "ad": {"status", "bid_amt", "delete"},
    "ad_extension": {"status"},  # 파워링크 확장소재: ON/OFF만 (내용 수정은 유형별 구조 확인 후)
    # 제외검색어: create는 naver_id=광고그룹·value=검색어, delete는 naver_id=제외검색어 ID·value=광고그룹 ID(삭제 경로에 필요)
    "negative_keyword": {"create", "delete"},
}
# shortcut: 입찰가·예산 한도는 가정값, 공식 문서 확인 후 수정
BID_RANGE = {"WEB_SITE": (70, 100_000), "SHOPPING": (50, 100_000)}
MIN_BUDGET = 50
EXCEL_HEADER = ["entity_type", "naver_id", "field", "value"]
# shortcut: 무료 플랜 CPU 한도 안에서 미리보기를 만들려고 정한 값(실측 아님), 큰 엑셀이 필요해지면 실측 후 올리거나 나눠 저장
MAX_JOB_ROWS = 1000


NAME_MAX = 30  # shortcut: 네이버 화면 기준 이름 길이(명세 미확인), 네이버가 거부하면 오류로 보인다
MAX_BUDGET = 1_000_000_000


def _parts(value: str, n: int) -> list[str]:
    return [x.strip() for x in (value.split("|") + [""] * n)[:n]]


def _money(v: str, label: str, lo: int, hi: int) -> str | None:
    if not v.isdigit() or int(v) % 10 or not lo <= int(v) <= hi:
        return f"{label}은 {lo:,}~{hi:,}원, 10원 단위"
    return None


def _validate_create(etype: str, nid: str, value: str, product: str | None) -> str | None:
    if etype == "campaign":
        name, budget = _parts(value, 2)
        if nid not in BID_RANGE:
            return "캠페인 유형은 WEB_SITE(파워링크) 또는 SHOPPING(쇼핑검색)"
        if not name or len(name) > NAME_MAX or value.count("|") > 1:
            return f"캠페인 이름은 1~{NAME_MAX}자, | 없이"
        return _money(budget, "일예산", MIN_BUDGET, MAX_BUDGET)
    name, channel, bid, budget = _parts(value, 4)
    if not name or len(name) > NAME_MAX or value.count("|") > 3:
        return f"광고그룹 이름은 1~{NAME_MAX}자, | 없이"
    if not channel.startswith("bsn-"):
        return "비즈채널 ID(bsn-...)가 필요합니다"
    lo, hi = BID_RANGE[product or "WEB_SITE"]
    return (bid and _money(bid, "입찰가", lo, hi)) or (budget and _money(budget, "하루예산", MIN_BUDGET, MAX_BUDGET)) or None


def _split_new_keyword(value: str) -> tuple[str, str]:
    kw, _, bid = value.partition("|")
    return kw.strip(), bid.strip()


def _validate(ch: dict) -> str | None:
    """ch: entity_type, field, value, before, product(WEB_SITE|SHOPPING, 키워드 등록은 광고그룹 유형), ad_type.
    product·ad_type이 없으면(실패 행 재실행) 그 검사는 건너뛴다 — 처음 미리보기에서 이미 통과한 행이다."""
    etype, field, value, before, product = ch["entity_type"], ch["field"], ch["value"], ch.get("before"), ch.get("product")
    if field not in ALLOWED_FIELDS[etype]:
        return f"{etype}에서 수정할 수 없는 항목: {field}"
    if etype == "negative_keyword":
        return None if value else "제외검색어를 입력하세요" if field == "create" else "광고그룹 ID가 필요합니다"
    if field == "delete":
        return None
    if field == "create" and etype in ("campaign", "adgroup"):
        return _validate_create(etype, ch["naver_id"], value, product)
    if field == "create":
        kw, bid = _split_new_keyword(value)
        if product and product != "WEB_SITE":
            return "키워드는 파워링크 광고그룹에만 등록할 수 있습니다"
        if not kw:
            return "키워드를 입력하세요"
        lo, hi = BID_RANGE["WEB_SITE"]
        if bid and (not bid.isdigit() or int(bid) % 10 or not lo <= int(bid) <= hi):
            return f"입찰가는 {lo:,}~{hi:,}원, 10원 단위 (비우면 {lo}원)"
        return None
    if field == "status":
        if value not in ("ON", "OFF"):
            return "상태는 ON 또는 OFF"
    else:
        if etype == "ad" and ch.get("ad_type") not in (None, "SHOPPING_PRODUCT_AD"):
            return "입찰가는 쇼핑 상품소재에만 설정할 수 있습니다"
        if not value.isdigit():
            return "숫자만 입력하세요"
        n = int(value)
        if n % 10:
            return "10원 단위로 입력하세요"
        if field == "daily_budget" and n < MIN_BUDGET:
            return f"일예산은 {MIN_BUDGET}원 이상"
        if field == "bid_amt":
            lo, hi = BID_RANGE["WEB_SITE" if etype == "keyword" else product or "SHOPPING"]
            if not lo <= n <= hi:
                return f"입찰가는 {lo:,}~{hi:,}원"
    if before is not None and str(before) == value:
        return "현재 값과 동일합니다"
    return None


def create_job(db: Session, m: Membership, changes: list[dict], source: str, parent_job_id: int | None = None) -> BulkJob:
    """changes: entity_type, naver_id, field, value, ad_account_id, before?, product?, ad_type?, invalid?(화면이 찾은 오류)"""
    if not changes:
        raise HTTPException(400, "변경 내용이 없습니다")
    if len(changes) > MAX_JOB_ROWS:
        raise HTTPException(400, f"한 번에 {MAX_JOB_ROWS:,}행까지입니다. 나눠서 올려 주세요")
    allowed = accessible_account_ids(db, m, edit=True)
    job = BulkJob(workspace_id=m.workspace_id, created_by=m.user_id, source=source, parent_job_id=parent_job_id)
    db.add(job)
    db.flush()
    seen = set()
    rows = []
    for ch in changes:
        etype, nid, field, value = (str(ch.get(k) or "").strip() for k in EXCEL_HEADER)
        ch = {**ch, "entity_type": etype, "naver_id": nid, "field": field, "value": value}
        before = ch.get("before")
        row = BulkJobRow(job_id=job.id, entity_type=etype, naver_id=nid, field=field, after=value, status="INVALID",
                         before=None if before is None else str(before)[:50])
        # 계정까지 같아야 중복: 캠페인 등록은 naver_id가 유형(WEB_SITE)이라 계정이 다르면 다른 대상이다
        key = (ch.get("ad_account_id"), nid, field, _split_new_keyword(value)[0] if field == "create" else "")
        if etype not in ALLOWED_FIELDS:
            row.error = f"알 수 없는 유형: {etype}"
        elif not nid or ch.get("ad_account_id") not in allowed:
            row.error = "대상이 없거나 수정 권한이 없습니다"  # 권한 없는 계정 이름을 노출하지 않는다
        elif key in seen:
            row.ad_account_id, row.error = ch["ad_account_id"], "같은 대상·항목이 중복되었습니다"
        else:
            seen.add(key)
            row.ad_account_id = ch["ad_account_id"]
            row.error = str(ch["invalid"])[:300] if ch.get("invalid") else _validate(ch)
            if row.error is None:
                row.status = "VALID"
        rows.append(row)
    db.add_all(rows)  # 행마다 INSERT를 따로 보내지 않게 한 번에(무료 플랜 요청당 쿼리 한도)
    db.commit()
    return job


def _client(db: Session, workspace_id: int, account_id: int):
    link = db.scalar(select(WorkspaceAdAccount).where(
        WorkspaceAdAccount.workspace_id == workspace_id, WorkspaceAdAccount.ad_account_id == account_id))
    return client_for(db.get(NaverApiCredential, link.credential_id), db.get(AdAccount, account_id).customer_id, decrypt)


CUT_OFF = "실행 중 연결이 끊겼습니다. 네이버에서 반영됐는지 확인하세요"
# 무료 플랜은 요청 1번에 D1 쿼리+외부 호출 50번까지라, 화면 실행은 몇 행씩 나눠 여러 번 요청한다(행마다 약 4번 + 고정 약 10번)
EXECUTE_CHUNK = 5


def execute_job(db: Session, m: Membership, job: BulkJob, limit: int | None = None) -> BulkJob:
    """VALID 행을 최대 limit개 실행. 남은 행이 있으면 RUNNING으로 두고 반환하니 다시 호출하면 이어서 실행한다
    (중간에 요청이 끊긴 RUNNING 작업도 같은 방법으로 이어 간다)."""
    if job.status not in ("PREVIEW", "RUNNING"):
        raise HTTPException(409, "이미 실행된 작업입니다")
    q = select(BulkJobRow).where(BulkJobRow.job_id == job.id, BulkJobRow.status == "VALID").order_by(BulkJobRow.id)
    rows = db.scalars(q.limit(limit) if limit else q).all()
    if not rows and job.status == "PREVIEW":
        raise HTTPException(400, "실행할 수 있는 행이 없습니다")
    job.status = "RUNNING"
    db.commit()
    allowed = accessible_account_ids(db, m, edit=True)  # 미리보기 이후 권한이 바뀌었을 수 있다
    clients = {}
    for row in rows:
        if row.ad_account_id not in allowed:
            row.status, row.error = "FAILED", "실행 시점에 수정 권한이 없습니다"
            continue
        # 같은 작업을 두 곳에서 동시에 실행해도(Claude가 겹쳐 부름 등) 한 행은 한 번만: 먼저 RUNNING으로 바꾼 요청만 실행
        if db.execute(update(BulkJobRow).where(BulkJobRow.id == row.id, BulkJobRow.status == "VALID")
                      .values(status="RUNNING")).rowcount != 1:
            continue
        db.commit()
        try:
            if row.ad_account_id not in clients:
                clients[row.ad_account_id] = _client(db, m.workspace_id, row.ad_account_id)
            client = clients[row.ad_account_id]
            if row.entity_type == "negative_keyword":
                if row.field == "create":
                    created = client.create_restricted_keyword(row.naver_id, row.after)
                else:
                    client.delete_restricted_keyword(row.after, row.naver_id)
            elif row.field == "create" and row.entity_type == "campaign":
                name, budget = _parts(row.after, 2)
                created = client.create_campaign(row.naver_id, name, int(budget))
            elif row.field == "create" and row.entity_type == "adgroup":
                name, channel, bid, budget = _parts(row.after, 4)
                created = client.create_adgroup(row.naver_id, name, channel, int(bid) if bid else None, int(budget) if budget else None)
            elif row.field == "create":
                kw, bid = _split_new_keyword(row.after)
                bid_amt = int(bid) if bid else BID_RANGE["WEB_SITE"][0]  # 비우면 최저 입찰가(사용자 요청 10/20)
                created = client.create_keyword(row.naver_id, kw, bid_amt, False)
            elif row.field == "delete" and row.entity_type in ("campaign", "adgroup"):
                client.delete_container(row.entity_type, row.naver_id)
            elif row.field == "delete":
                (client.delete_keyword if row.entity_type == "keyword" else client.delete_ad)(row.naver_id)
            else:  # 예약 작업은 미리보기와 실행 사이가 길어 그 사이 바뀐 값을 덮어쓰지 않는다
                client.update(row.entity_type, row.naver_id, row.field, row.after, expect=row.before if job.scheduled_at else None)
        except NaverApiError as e:
            row.status, row.error = "FAILED", str(e)[:300]
            continue
        audit = dict(action="UPDATE", naver_id=row.naver_id, field=row.field, before=row.before, after=row.after)
        if row.entity_type == "negative_keyword" and row.field == "create":
            audit.update(action="CREATE", naver_id=created["nccAdgroupRestrictKwdId"], field="negative_keyword", after=row.after)
        elif row.entity_type == "negative_keyword":
            audit.update(action="DELETE", field="negative_keyword", after=None)
        elif row.field == "create" and row.entity_type in ("campaign", "adgroup"):
            nid_key = "nccCampaignId" if row.entity_type == "campaign" else "nccAdgroupId"
            audit.update(action="CREATE", naver_id=created[nid_key], field=row.entity_type, after=_parts(row.after, 1)[0])
        elif row.field == "create":
            audit.update(action="CREATE", naver_id=created["nccKeywordId"], field="keyword", after=kw)
        elif row.field == "delete":
            if row.entity_type == "keyword":
                db.execute(delete(AutoBidRule).where(AutoBidRule.workspace_id == m.workspace_id, AutoBidRule.keyword_naver_id == row.naver_id))
            audit.update(action="DELETE", field=row.entity_type, after=None)
        row.status = "SUCCESS"
        db.add(AuditLog(workspace_id=m.workspace_id, user_id=m.user_id, ad_account_id=row.ad_account_id, job_id=job.id,
                        entity_type=row.entity_type, **audit))
        db.commit()  # 행 단위 커밋: 중간에 죽어도 성공한 변경과 감사 로그는 남는다
    db.commit()
    results = set(db.scalars(select(BulkJobRow.status).where(BulkJobRow.job_id == job.id, BulkJobRow.status != "INVALID")))
    if "VALID" in results:
        return job
    if "RUNNING" in results:  # 실행 중 요청이 끊긴 행: 네이버 반영 여부를 알 수 없어 다시 실행하지 않는다
        db.execute(update(BulkJobRow).where(BulkJobRow.job_id == job.id, BulkJobRow.status == "RUNNING")
                   .values(status="FAILED", error=CUT_OFF))
        results = (results - {"RUNNING"}) | {"FAILED"}
    job.status = "DONE" if results == {"SUCCESS"} else "FAILED" if "SUCCESS" not in results else "PARTIAL"
    job.executed_at = now()
    db.commit()
    return job


def retry_failed(db: Session, m: Membership, job: BulkJob) -> BulkJob:
    failed = db.scalars(select(BulkJobRow).where(BulkJobRow.job_id == job.id, BulkJobRow.status == "FAILED")).all()
    if not failed:
        raise HTTPException(400, "재실행할 실패 행이 없습니다")
    # 연결이 끊긴 행은 네이버에 이미 반영됐을 수 있어(등록이면 두 번 생김) 재실행하지 않는 행으로 넘긴다
    changes = [{"entity_type": r.entity_type, "naver_id": r.naver_id, "field": r.field, "value": r.after,
                "ad_account_id": r.ad_account_id, "before": r.before,
                **({"invalid": "반영 여부를 알 수 없어 다시 실행하지 않습니다. 네이버에서 확인 후 새로 만드세요"} if r.error == CUT_OFF else {})}
               for r in failed]
    return create_job(db, m, changes, "RETRY", parent_job_id=job.id)


def _group_ads(client, adgroup_ids) -> dict[str, set[str]]:
    return {g: {a["nccAdId"] for a in client.list_ads(g)} for g in set(adgroup_ids)}


def _rank(serp: list[str], mine: set[str]) -> int | None:
    """검색 화면 소재 ID 목록에서 내 광고그룹 소재가 처음 나오는 순위. 없으면 None(노출 안 됨)."""
    return next((i for i, x in enumerate(serp, 1) if x in mine), None)


def next_bid(cur: int, rank: int | None, target: int, max_bid: int) -> int:
    """목표보다 아래(또는 노출 안 됨)면 20% 올리고, 위면 10% 내린다(최소 10원). 내릴 때 덜 움직여 순위가 출렁이지 않게."""
    lo = BID_RANGE["WEB_SITE"][0]
    if rank == target:
        return cur
    if rank is None or rank > target:
        return max(lo, min(cur + max(10, cur // 5 // 10 * 10), max_bid // 10 * 10))
    return max(lo, cur - max(10, cur // 10 // 10 * 10))


def check_ranks(db: Session, workspace_id: int, refs) -> list[dict]:
    """지금 순위: 키워드별 PC·모바일 검색 화면에서 내 광고가 몇 위인지. refs는 권한 확인을 마친 KeywordRef."""
    out = []
    by_account: dict[int, list] = {}
    for k in refs:
        by_account.setdefault(k.ad_account_id, []).append(k)
    for account_id, ks in by_account.items():
        try:
            client = _client(db, workspace_id, account_id)
            kws = {k["nccKeywordId"]: k for k in client.get_many("keywords", [k.naver_id for k in ks])}
            ads = _group_ads(client, [k["nccAdgroupId"] for k in kws.values()])
        except NaverApiError as e:
            out += [{"naver_id": k.naver_id, "error": f"네이버 조회 실패: {e}"[:200]} for k in ks]
            continue
        for k in ks:
            kw = kws.get(k.naver_id)
            row = {"naver_id": k.naver_id}
            if kw is None:
                row["error"] = "네이버에 키워드가 없습니다"
            else:
                try:
                    for d in ("PC", "MOBILE"):
                        row[d] = _rank(client.serp_ads(kw["keyword"], d), ads[kw["nccAdgroupId"]])
                except NaverApiError as e:
                    row["error"] = str(e)[:200]
            out.append(row)
    return out


AUTO_BID_DAILY_MAX = 24  # 키워드당 하루 자동 변경 횟수(크론은 하루 144번). 순위가 출렁여 입찰가가 오르내리기만 반복하는 것을 막는다


def _auto_bid_changes_today(db: Session, workspace_id: int) -> dict[str, int]:
    day = datetime.now(KST).replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc)
    q = (select(AuditLog.naver_id, func.count()).join(BulkJob, BulkJob.id == AuditLog.job_id)
         .where(BulkJob.workspace_id == workspace_id, BulkJob.source == "AUTO_BID", AuditLog.created_at >= day)
         .group_by(AuditLog.naver_id))
    return dict(db.execute(q).all())


def run_auto_bids(db: Session, workspace_id: int, rule_ids: list[int] | None = None) -> list[BulkJob]:
    """자동입찰 규칙 실행: 검색 화면에서 지금 순위를 읽어 목표보다 아래면 올리고 위면 내린다(next_bid, 최대 입찰가 이하).
    바꿀 키워드는 만든 사람별 대량 작업으로 실행한다. 권한 확인·네이버 반영 검증·감사 로그는 일반 대량 작업과 같다."""
    s = select(AutoBidRule).where(AutoBidRule.workspace_id == workspace_id, AutoBidRule.enabled.is_(True))
    if rule_ids is not None:
        s = s.where(AutoBidRule.id.in_(rule_ids))
    by_user: dict[int, list[AutoBidRule]] = {}
    for r in db.scalars(s):
        by_user.setdefault(r.created_by, []).append(r)
    jobs = []
    changed_today = _auto_bid_changes_today(db, workspace_id) if by_user else {}
    for user_id, rules in by_user.items():
        m = db.scalar(select(Membership).where(Membership.workspace_id == workspace_id, Membership.user_id == user_id))
        allowed = accessible_account_ids(db, m, edit=True) if m else set()
        by_account: dict[int, list[AutoBidRule]] = {}
        for r in rules:
            r.last_run_at = now()
            if r.ad_account_id not in allowed:
                r.last_result = "규칙을 만든 사람에게 이 계정 수정 권한이 없습니다"
            else:
                by_account.setdefault(r.ad_account_id, []).append(r)
        changes, rule_by_nid = [], {}
        for account_id, rs in by_account.items():
            try:
                client = _client(db, workspace_id, account_id)
                current = {k["nccKeywordId"]: k for k in client.get_many("keywords", [r.keyword_naver_id for r in rs])}
                gids = list({k["nccAdgroupId"] for k in current.values()})
                groups = {g["nccAdgroupId"]: g for g in client.get_many("adgroups", gids)} if gids else {}
                ads = _group_ads(client, gids)
            except NaverApiError as e:
                for r in rs:
                    r.last_result = f"네이버 조회 실패: {e}"[:300]
                continue
            for r in rs:
                kw = current.get(r.keyword_naver_id)
                if kw is None:
                    r.last_result = "네이버에 키워드가 없습니다 (삭제되었으면 규칙을 해제하세요)"
                    continue
                grp = groups.get(kw["nccAdgroupId"], {})
                # 꺼져 있으면 노출이 안 돼 입찰가만 계속 오르므로 건너뛴다. shortcut: 캠페인 OFF는 안 봄, 필요하면 캠페인도 조회
                if kw.get("userLock") or grp.get("userLock"):
                    r.last_result = "키워드나 광고그룹이 OFF라 건너뜀"
                    continue
                if changed_today.get(r.keyword_naver_id, 0) >= AUTO_BID_DAILY_MAX:
                    r.last_result = f"오늘 자동 변경 {AUTO_BID_DAILY_MAX}번을 채워 내일 다시 조정"
                    continue
                try:
                    rank = _rank(client.serp_ads(kw["keyword"], r.device), ads[kw["nccAdgroupId"]])
                except NaverApiError as e:
                    r.last_result = f"순위 확인 실패: {e}"[:300]
                    continue
                cur = grp.get("bidAmt", kw["bidAmt"]) if kw.get("useGroupBidAmt") else kw["bidAmt"]
                target = next_bid(cur, rank, r.target_rank, r.max_bid)
                r.last_result = f"지금 {rank}위" if rank else "지금 노출 안 됨"
                r.last_result += f" (목표 {r.target_rank}위)"
                if target == cur:
                    r.last_result += " · 변경 없음" + (" (최대 입찰가)" if rank != r.target_rank and cur >= r.max_bid else "")
                    continue
                r.last_result += f" · {cur:,}원 → {target:,}원"
                changes.append({"entity_type": "keyword", "naver_id": r.keyword_naver_id, "field": "bid_amt", "value": str(target),
                                "ad_account_id": account_id, "before": cur, "product": "WEB_SITE"})
                rule_by_nid[r.keyword_naver_id] = r
        db.commit()
        if not changes:
            continue
        job = execute_job(db, m, create_job(db, m, changes, "AUTO_BID"))
        for row in db.scalars(select(BulkJobRow).where(BulkJobRow.job_id == job.id, BulkJobRow.status != "SUCCESS")):
            rule_by_nid[row.naver_id].last_result = f"{rule_by_nid[row.naver_id].last_result} 실패: {row.error}"[:300]
        db.commit()
        jobs.append(job)
    return jobs


def run_all_auto_bids(db: Session):
    """크론용: 규칙이 있는 모든 조직. 한 조직이 실패해도 나머지는 돈다."""
    for ws_id in set(db.scalars(select(AutoBidRule.workspace_id).where(AutoBidRule.enabled.is_(True)))):
        try:
            run_auto_bids(db, ws_id)
        except Exception as e:  # noqa: BLE001
            db.rollback()
            print(f"자동입찰 실패 workspace={ws_id}: {e}")


def _check_edit(db: Session, m: Membership, account_id: int):
    if account_id not in accessible_account_ids(db, m, edit=True):
        raise HTTPException(403, "이 광고 계정을 수정할 권한이 없습니다")


def create_extension(db: Session, m: Membership, account_id: int, adgroup_naver_id: str, ext_type: str, content) -> dict:
    """확장소재 생성은 행 단위 미리보기 없이 바로 실행하되, 권한 확인과 감사 로그는 같은 규칙을 따른다. 화면 보관본용 행을 돌려준다."""
    _check_edit(db, m, account_id)
    try:
        e = _client(db, m.workspace_id, account_id).create_ad_extension(adgroup_naver_id, ext_type, content)
    except NaverApiError as err:
        raise HTTPException(400, f"네이버 확장소재 생성 실패: {err}")
    row = extension_row(account_id, {"type": ext_type, **e})
    if not e.get("adExtension"):
        row["payload"] = content
    db.add(AuditLog(workspace_id=m.workspace_id, user_id=m.user_id, ad_account_id=account_id, action="CREATE",
                    entity_type="ad_extension", naver_id=row["naver_id"], field="ext_type", after=row["ext_type"]))
    db.commit()
    return row


def create_ad(db: Session, m: Membership, account_id: int, adgroup_naver_id: str, ad_type: str, headlines: list[str],
              descriptions: list[str], pc_url: str, mobile_url: str) -> dict:
    """소재 등록(OFF로 생성). 확장소재처럼 행 단위 미리보기 없이 바로 실행하고, 권한 확인·감사 로그는 같다.
    RSA_AD(반응형)는 제목·설명을 assets로, TEXT_45는 첫 제목·설명만 ad에 넣는다. 글자 수·개수 검사는 네이버에 맡긴다."""
    _check_edit(db, m, account_id)
    headlines, descriptions = [h.strip() for h in headlines if h.strip()], [d.strip() for d in descriptions if d.strip()]
    if not headlines or not descriptions or not pc_url.startswith("http"):
        raise HTTPException(400, "제목·설명을 1개 이상, 연결 URL은 http로 시작하게 입력하세요")
    ad = {"pc": {"final": pc_url}, "mobile": {"final": mobile_url or pc_url}}
    assets = None
    if ad_type == "RSA_AD":
        assets = ([{"linkType": "HEADLINE", "assetData": {"text": h}} for h in headlines]
                  + [{"linkType": "DESCRIPTION", "assetData": {"text": d}} for d in descriptions])
    else:
        ad |= {"headline": headlines[0], "description": descriptions[0]}
    try:
        a = _client(db, m.workspace_id, account_id).create_ad(adgroup_naver_id, ad_type, ad, assets)
    except NaverApiError as err:
        raise HTTPException(400, f"네이버 소재 등록 실패: {err}")
    row = ad_row(account_id, a)
    db.add(AuditLog(workspace_id=m.workspace_id, user_id=m.user_id, ad_account_id=account_id, action="CREATE",
                    entity_type="ad", naver_id=row["naver_id"], field="ad_type", after=ad_type))
    db.commit()
    return row


def delete_extension(db: Session, m: Membership, account_id: int, ext_naver_id: str, owner_naver_id: str, ext_type: str | None) -> bool:
    _check_edit(db, m, account_id)
    client = _client(db, m.workspace_id, account_id)
    gone = False
    try:
        client.delete_ad_extension(ext_naver_id)
    except NaverApiError as err:
        # 네이버에서 이미 지워진 확장소재는 404가 난다(보관본에는 남아 있을 수 있다): 소유자 목록으로 실제 존재 여부를 확인
        try:
            gone = all(e["nccAdExtensionId"] != ext_naver_id for e in client.list_ad_extensions(owner_naver_id))
        except NaverApiError:
            gone = False
        if not gone:
            raise HTTPException(400, f"네이버 확장소재 삭제 실패: {err}")
    db.add(AuditLog(workspace_id=m.workspace_id, user_id=m.user_id, ad_account_id=account_id, action="DELETE",
                    entity_type="ad_extension", naver_id=ext_naver_id, field="ext_type", before=ext_type,
                    after="네이버에 이미 없음" if gone else None))
    db.commit()
    return gone


def parse_excel(data: bytes) -> list[dict]:
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True).active
    except Exception:
        raise HTTPException(400, "엑셀(.xlsx) 파일을 읽을 수 없습니다")
    rows = list(ws.iter_rows(values_only=True))
    if not rows or [str(c or "").strip() for c in rows[0][:4]] != EXCEL_HEADER:
        raise HTTPException(400, f"첫 행은 {EXCEL_HEADER} 이어야 합니다")
    return [dict(zip(EXCEL_HEADER, ("" if c is None else str(c) for c in r[:4]))) for r in rows[1:] if any(r[:4])]


def excel_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "bulk"
    ws.append(EXCEL_HEADER)
    ws.append(["campaign", "cmp-...", "status", "OFF"])
    ws.append(["adgroup", "grp-...", "bid_amt", "500"])
    ws.append(["ad", "ad-...(쇼핑 상품소재)", "bid_amt", "300"])
    ws.append(["keyword", "grp-...(광고그룹)", "create", "여름 원피스|500"])
    ws.append(["keyword", "kwd-...", "delete", ""])
    guide = wb.create_sheet("guide")
    for line in [["entity_type", "campaign | adgroup | keyword(파워링크) | ad | ad_extension(파워링크 확장소재)"],
                 ["field", "status(ON/OFF), daily_budget(캠페인), bid_amt(광고그룹·키워드·쇼핑 상품소재), 확장소재는 status만"],
                 ["value", "10원 단위 숫자 또는 ON/OFF"],
                 ["키워드 등록", "entity_type=keyword, naver_id=광고그룹 ID(grp-...), field=create, value=키워드|입찰가 (입찰가 비우면 70원)"],
                 ["키워드 삭제", "entity_type=keyword, naver_id=키워드 ID(kwd-...), field=delete, value 비움"],
                 ["소재 삭제", "entity_type=ad, naver_id=소재 ID(nad-...), field=delete, value 비움"],
                 ["제외검색어 등록", "entity_type=negative_keyword, naver_id=광고그룹 ID(grp-...), field=create, value=제외할 검색어"],
                 ["제외검색어 삭제", "entity_type=negative_keyword, naver_id=제외검색어 ID, field=delete, value=광고그룹 ID(grp-...)"]]:
        guide.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- 예약 실행 ----
def rotate_credentials(db: Session):
    """크론용: 예전 열쇠(APP_SECRET)로 암호화된 네이버 비밀키를 CRED_KEY로 다시 암호화한다."""
    for cred in db.scalars(select(NaverApiCredential)):
        new = rotated(cred.secret_key_enc)
        if new:
            cred.secret_key_enc = new
    db.commit()


def run_scheduled_jobs(db: Session):
    """크론용: 예약 시각이 지난 작업 하나를 EXECUTE_CHUNK행 실행(남으면 다음 크론에서 이어서). 실행은 작업을 만든 사람 권한으로.
    shortcut: 크론 1번에 5행이라 큰 작업은 오래 걸린다, 예약 작업이 커지면 한 번에 더 실행한다"""
    job = db.scalar(select(BulkJob).where(BulkJob.scheduled_at <= now(), BulkJob.status.in_(("PREVIEW", "RUNNING")))
                    .order_by(BulkJob.scheduled_at).limit(1))
    if job is None:
        return
    m = db.scalar(select(Membership).where(Membership.workspace_id == job.workspace_id, Membership.user_id == job.created_by))
    try:
        if m is None:
            raise HTTPException(403, "작업을 만든 사람이 조직에 없습니다")
        execute_job(db, m, job, EXECUTE_CHUNK)
    except Exception as e:  # noqa: BLE001 — 권한·행 없음·예상 못 한 오류 모두 끝내야 매 크론 같은 작업에 막히지 않는다
        db.rollback()
        job.status, job.executed_at = "FAILED", now()
        db.add(AuditLog(workspace_id=job.workspace_id, user_id=job.created_by, job_id=job.id, action="JOB_SCHEDULE_FAIL",
                        after=str(getattr(e, "detail", e))[:50]))
        db.commit()


# ---- 계정 목표 ROAS ----
# MVP 정책: ROAS는 보고서·분석·추천용이고 ROAS로 입찰가·상태를 자동으로 바꾸지 않는다.
# 코드는 남겨 두되 ROAS_AUTOMATION=1일 때만 실행(크론·지금 실행 모두). 정책이 바뀌기 전엔 켜지 말 것
def roas_automation_on() -> bool:
    return os.getenv("ROAS_AUTOMATION") == "1"


GOAL_RUN_HOUR = 9  # 어제 통계가 다 쌓인 뒤 하루 한 번(한국 시간)
GOAL_MIN_CLICKS = 10  # 이보다 클릭이 적은 키워드는 평균이 흔들려 그대로 둔다


def goal_bid(cur: int, st: dict, target: int, max_bid: int, off_cost: int | None, allow_raise: bool) -> int | str | None:
    """키워드 최근 7일 성과 st로 할 일: "OFF", 새 입찰가, None(그대로).
    ROAS가 목표 아래면 -10%, 목표보다 10% 넘게 높으면 +10%(계정 ROAS가 목표 이상일 때만, 최대 입찰가까지).
    shortcut: 고정 10%씩, 성과가 쌓이면 목표와의 차이만큼 움직이게 바꾼다"""
    clk, cost, conv, rev = st.get("clkCnt", 0), st.get("salesAmt", 0), st.get("ccnt") or 0, st.get("convAmt") or 0
    if off_cost and cost >= off_cost and not conv:
        return "OFF"
    if clk < GOAL_MIN_CLICKS or not conv or not cost:
        return None
    roas, step = rev * 100 / cost, max(10, cur // 10 // 10 * 10)
    if roas < target:
        new = max(BID_RANGE["WEB_SITE"][0], cur - step)
    elif roas >= target * 1.1 and allow_raise:
        new = max(cur, min(cur + step, max_bid // 10 * 10))
    else:
        return None
    return new if new != cur else None


def _goal_result(g: AccountGoal, st: dict) -> str:
    roas = "없음" if st["roas"] is None else f"{st['roas']:,}%"
    out = f"{st['until'][5:]}까지 7일 계정 ROAS {roas} (목표 {g.roas_target:,}%)"
    if st["roas"] is None or st["roas"] < g.roas_target:
        out += " · 목표 미달이라 내리기만"
    if g.off_cost and not st.get("conv"):
        out += " · 계정 전환 0건이라 무전환 OFF 안 함(전환 추적 확인)"
    out += f" · 키워드 {st['checked']}개 확인, 입찰가 조정 {st['bid']}개, OFF {st['off']}개"
    return out + (f" · 진행 중(광고그룹 {len(st['groups'])}개 남음)" if st["groups"] or st.get("job") else "")


def run_goal_step(db: Session, g: AccountGoal, start: bool = False) -> bool:
    """한 단계만 실행한다(무료 플랜은 실행 1번에 외부 호출 50번까지): 진행 중 작업 이어 실행 → 남은 광고그룹 1개 조정 → 오늘 첫 시작.
    start=True면 시간과 상관없이 새로 시작한다(지금 실행). 할 일이 없으면 False.
    shortcut: 크론 1번에 1단계라 광고그룹이 수백 개면 하루에 다 못 돈다, 광고주가 많아지면 단계 수를 늘린다"""
    today = datetime.now(KST)
    st = json.loads(g.state) if g.state else None
    if st is None and not start and (today.hour < GOAL_RUN_HOUR or g.run_date == today.date().isoformat()):
        return False
    g.last_run_at = now()
    m = db.scalar(select(Membership).where(Membership.workspace_id == g.workspace_id, Membership.user_id == g.created_by))
    if not m or g.ad_account_id not in accessible_account_ids(db, m, edit=True):
        g.last_result, g.state, g.run_date = "목표를 저장한 사람에게 이 계정 수정 권한이 없습니다", None, today.date().isoformat()
        db.commit()
        return True
    try:
        client = _client(db, g.workspace_id, g.ad_account_id)
        job = st and st.get("job") and db.get(BulkJob, st["job"])
        if st and not (job and job.status == "RUNNING"):
            st["job"] = None
        if st and st["job"]:
            execute_job(db, m, job, EXECUTE_CHUNK)
            st["job"] = job.id if job.status == "RUNNING" else None
        elif st is None:
            until = today.date() - timedelta(days=1)
            since = until - timedelta(days=6)
            camps = [c["nccCampaignId"] for c in client.list_campaigns() if c["campaignTp"] == "WEB_SITE" and not c.get("userLock")]
            s = client.entity_stats(camps, since, until).values() if camps else []
            cost, rev, conv = (sum(x.get(k) or 0 for x in s) for k in ("salesAmt", "convAmt", "ccnt"))
            groups = [x["nccAdgroupId"] for c in camps for x in client.list_adgroups(c) if not x.get("userLock")]
            st = {"groups": groups, "roas": round(rev * 100 / cost) if cost else None, "conv": conv, "since": since.isoformat(),
                  "until": until.isoformat(), "checked": 0, "bid": 0, "off": 0, "job": None}
            g.run_date = today.date().isoformat()
        elif st["groups"]:
            gid = st["groups"].pop(0)
            grp = (client.get_many("adgroups", [gid]) or [{}])[0]
            # 순위 자동입찰이 걸린 키워드는 그 규칙에 맡긴다(둘이 번갈아 올리고 내리지 않게)
            ruled = set(db.scalars(select(AutoBidRule.keyword_naver_id).where(AutoBidRule.workspace_id == g.workspace_id,
                                                                              AutoBidRule.enabled.is_(True))))
            kws = [k for k in client.list_keywords(gid) if not k.get("userLock") and k["nccKeywordId"] not in ruled]
            stats = client.entity_stats([k["nccKeywordId"] for k in kws], date.fromisoformat(st["since"]),
                                        date.fromisoformat(st["until"])) if kws else {}
            allow = st["roas"] is not None and st["roas"] >= g.roas_target
            # 계정 전체 전환이 0이면 전환 추적이 안 되는 것일 수 있어 무전환 OFF를 하지 않는다
            off_cost = g.off_cost if st.get("conv") else None
            changes = []
            for k in kws:
                cur = grp.get("bidAmt", k["bidAmt"]) if k.get("useGroupBidAmt") else k["bidAmt"]
                act = goal_bid(cur, stats.get(k["nccKeywordId"], {}), g.roas_target, g.max_bid, off_cost, allow)
                st["checked"] += 1
                ch = {"entity_type": "keyword", "naver_id": k["nccKeywordId"], "ad_account_id": g.ad_account_id, "product": "WEB_SITE"}
                if act == "OFF":
                    changes.append({**ch, "field": "status", "value": "OFF", "before": "ON"})
                    st["off"] += 1
                elif act:
                    changes.append({**ch, "field": "bid_amt", "value": str(act), "before": cur})
                    st["bid"] += 1
            if changes:
                job = execute_job(db, m, create_job(db, m, changes, "ROAS_GOAL"), EXECUTE_CHUNK)
                st["job"] = job.id if job.status == "RUNNING" else None
    except NaverApiError as e:
        g.last_result, g.state, g.run_date = f"네이버 조회 실패: {e}"[:300], None, today.date().isoformat()
        db.commit()
        return True
    g.last_result = _goal_result(g, st)[:300]
    g.state = json.dumps(st) if st["groups"] or st.get("job") else None
    db.commit()
    return True


def run_all_goals(db: Session):
    """크론용: 할 일이 있는 첫 목표 하나만 한 단계 진행(무료 플랜 호출 한도). 실패해도 다음 크론에서 다른 목표로 넘어간다."""
    if not roas_automation_on():
        return
    for g in db.scalars(select(AccountGoal).where(AccountGoal.enabled.is_(True)).order_by(AccountGoal.id)):
        try:
            if run_goal_step(db, g):
                return
        except Exception as e:  # noqa: BLE001
            db.rollback()
            print(f"목표 ROAS 실패 goal={g.id}: {e}")
